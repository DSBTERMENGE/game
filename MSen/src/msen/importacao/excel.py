"""Leitor somente leitura para planilhas de resultados em formato XLSX."""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from unicodedata import normalize
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from msen.modelos import ErroValidacaoResultado, ResultadoEntrada


CAMPOS_OBRIGATORIOS = (
    "Concurso", "Data", "bola 1", "bola 2", "bola 3", "bola 4", "bola 5", "bola 6",
)


class ErroLeituraExcel(ValueError):
    """Erro identificável de estrutura ou conteúdo da planilha."""

    def __init__(self, codigo: str, mensagem: str, *, linha: int | None = None,
                 campo: str | None = None, valor: Any = None) -> None:
        super().__init__(mensagem)
        self.codigo = codigo
        self.linha = linha
        self.campo = campo
        self.valor = valor


@dataclass(frozen=True, slots=True)
class LinhaExcel:
    """Resultado da leitura tolerante de uma linha não vazia da planilha."""

    numero_linha: int
    numero_ap: object
    resultado: ResultadoEntrada | None = None
    erro: ErroLeituraExcel | None = None


def _nome_normalizado(valor: object) -> str:
    if not isinstance(valor, str):
        return ""
    sem_acentos = normalize("NFKD", valor).encode("ascii", "ignore").decode("ascii")
    return " ".join(sem_acentos.casefold().split())


def localizar_cabecalho(linhas: Iterable[Sequence[object]]) -> tuple[int, dict[str, int]]:
    """Retorna o número da linha e os índices das oito colunas obrigatórias."""

    exigidos = {_nome_normalizado(nome): nome for nome in CAMPOS_OBRIGATORIOS}
    melhor_linha: int | None = None
    melhor_encontro: set[str] = set()
    for numero_linha, linha in enumerate(linhas, start=1):
        encontrados: dict[str, int] = {}
        for indice, valor in enumerate(linha):
            nome = _nome_normalizado(valor)
            if nome in exigidos and nome not in encontrados:
                encontrados[nome] = indice
        if len(encontrados) == len(exigidos):
            return numero_linha, {
                original: encontrados[normalizado]
                for normalizado, original in exigidos.items()
            }
        if len(encontrados) > len(melhor_encontro):
            melhor_linha = numero_linha
            melhor_encontro = set(encontrados)

    ausentes = tuple(
        nome for nome in CAMPOS_OBRIGATORIOS if _nome_normalizado(nome) not in melhor_encontro
    )
    raise ErroLeituraExcel(
        "CABECALHO_NAO_ENCONTRADO",
        "Não foi encontrado um cabeçalho com todas as colunas obrigatórias.",
        linha=melhor_linha, campo="cabecalho", valor=ausentes,
    )


def _linha_vazia(linha: Sequence[object]) -> bool:
    return all(valor is None or (isinstance(valor, str) and not valor.strip()) for valor in linha)


def _valor_coluna(linha: Sequence[object], colunas: dict[str, int], nome: str) -> object:
    indice = colunas[nome]
    return linha[indice] if indice < len(linha) else None


def _normalizar_linha(linha: Sequence[object], colunas: dict[str, int],
                      numero_linha: int) -> ResultadoEntrada:
    data = _valor_coluna(linha, colunas, "Data")
    if isinstance(data, datetime):
        data = data.date()
    try:
        return ResultadoEntrada.criar(
            numero_ap=_valor_coluna(linha, colunas, "Concurso"), data_resultado=data,
            resultados=[_valor_coluna(linha, colunas, f"bola {indice}") for indice in range(1, 7)],
        )
    except ErroValidacaoResultado as erro:
        raise ErroLeituraExcel(
            "LINHA_INVALIDA", f"A linha {numero_linha} contém dados inválidos: {erro}",
            linha=numero_linha, campo=erro.campo, valor=erro.valor,
        ) from erro


def ler_linhas_excel(caminho: str | Path) -> tuple[LinhaExcel, ...]:
    """Lê todas as linhas e preserva erros de conteúdo para posterior relatório."""

    origem = Path(caminho)
    if origem.suffix.casefold() != ".xlsx":
        raise ErroLeituraExcel("FORMATO_NAO_SUPORTADO", "A fonte deve ser uma planilha no formato .xlsx.", campo="arquivo", valor=str(origem))
    if not origem.is_file():
        raise ErroLeituraExcel("ARQUIVO_NAO_ENCONTRADO", "A planilha informada não foi encontrada.", campo="arquivo", valor=str(origem))

    try:
        pasta = load_workbook(origem, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError):
        raise ErroLeituraExcel(
            "ARQUIVO_INVALIDO",
            "O arquivo enviado não é uma planilha XLSX válida.",
            campo="arquivo",
            valor=origem.name,
        ) from None
    try:
        folha = pasta.worksheets[0]
        numero_cabecalho, colunas = localizar_cabecalho(folha.iter_rows(values_only=True))
        linhas_lidas: list[LinhaExcel] = []
        for numero_linha, linha in enumerate(
            folha.iter_rows(min_row=numero_cabecalho + 1, values_only=True),
            start=numero_cabecalho + 1,
        ):
            if _linha_vazia(linha):
                continue
            numero_ap = _valor_coluna(linha, colunas, "Concurso")
            try:
                resultado = _normalizar_linha(linha, colunas, numero_linha)
            except ErroLeituraExcel as erro:
                linhas_lidas.append(LinhaExcel(numero_linha, numero_ap, erro=erro))
            else:
                linhas_lidas.append(LinhaExcel(numero_linha, numero_ap, resultado=resultado))
        return tuple(linhas_lidas)
    finally:
        pasta.close()


def ler_resultados_excel(caminho: str | Path) -> tuple[ResultadoEntrada, ...]:
    """Lê somente resultados válidos e interrompe diante da primeira linha inválida."""

    resultados: list[ResultadoEntrada] = []
    for linha in ler_linhas_excel(caminho):
        if linha.erro is not None:
            raise linha.erro
        if linha.resultado is not None:
            resultados.append(linha.resultado)
    return tuple(resultados)
